"""Testes dos endpoints de relatórios e exportação (PDF/CSV/XLSX)."""
import io
from openpyxl import load_workbook


def test_pdf_export_valid(client, gerente_login, produto):
    resp = client.get('/api/relatorio/pdf?tipo=produtos&dias=30')
    assert resp.status_code == 200
    assert resp.content_type.startswith('application/pdf')
    assert resp.headers.get('Content-Disposition', '').endswith('.pdf')
    assert resp.data[:5] == b'%PDF-'


def test_csv_export_valid(client, gerente_login, produto):
    resp = client.get('/api/relatorio/csv?tipo=produtos&dias=30')
    assert resp.status_code == 200
    assert 'text/csv' in resp.content_type
    text = resp.get_data(as_text=True)
    assert 'Produto' in text.splitlines()[0]


def test_xlsx_export_valid(client, gerente_login, produto):
    resp = client.get('/api/relatorio/xlsx?tipo=produtos&dias=30')
    assert resp.status_code == 200
    assert 'spreadsheetml' in resp.content_type
    assert resp.headers.get('Content-Disposition', '').endswith('.xlsx')
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert 'Produto' in headers
    nomes = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert 'Teste Produto' in nomes


def test_relatorio_tipo_invalido(client, gerente_login):
    resp = client.get('/api/relatorio/xlsx?tipo=inexistente')
    assert resp.status_code == 400
    assert resp.get_json()['ok'] is False


def test_relatorio_vendas_requer_login(client):
    # Sem sessão (não é requisito JSON) -> redireciona para a página de login
    resp = client.get('/api/relatorio/pdf?tipo=vendas')
    assert resp.status_code == 302
    assert '/login' in resp.headers.get('Location', '')


def test_resumo_diario(client, gerente_login):
    resp = client.get('/api/resumo')
    assert resp.status_code == 200
    dados = resp.get_json()
    assert 'vendas_hoje' in dados and 'compras_hoje' in dados