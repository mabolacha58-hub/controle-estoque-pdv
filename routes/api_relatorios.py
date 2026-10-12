import io
from datetime import datetime, date
from flask import Blueprint, request, jsonify, send_file
from database import moeda
from auth import login_required
from utils.pdf import gerar_pdf_relatorio
from services.relatorio_service import (resumo_diario, vendas_por_periodo, lucro_por_periodo,
    vendas_por_categoria, PDF_FUNCTIONS)

api_rel_bp = Blueprint('api_relatorios', __name__)

@api_rel_bp.route('/api/resumo')
@login_required
def api_resumo():
    return jsonify(resumo_diario(date.today().strftime('%Y-%m-%d')))

@api_rel_bp.route('/api/relatorio/vendas')
@login_required
def api_relatorio_vendas():
    dias = request.args.get('dias', 30, type=int)
    return jsonify(vendas_por_periodo(dias))

@api_rel_bp.route('/api/relatorio/lucro')
@login_required
def api_relatorio_lucro():
    dias = request.args.get('dias', 30, type=int)
    return jsonify(lucro_por_periodo(dias))

@api_rel_bp.route('/api/relatorio/vendas/categoria')
@login_required
def api_relatorio_vendas_categoria():
    dias = request.args.get('dias', 30, type=int)
    return jsonify(vendas_por_categoria(dias))

@api_rel_bp.route('/api/relatorio/csv')
@login_required
def api_relatorio_csv():
    tipo = request.args.get('tipo','vendas')
    dias = request.args.get('dias', 30, type=int)
    fn = PDF_FUNCTIONS.get(tipo)
    if not fn:
        return jsonify({'ok': False, 'erro': 'Tipo inválido'}), 400
    args = (dias,) if tipo in ('vendas','movimentacoes','lucro') else ()
    headers, data = fn(*args)
    import csv
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(headers)
    for row in data:
        writer.writerow(row)
    csv_bytes = io.BytesIO(output.getvalue().encode('utf-8-sig'))
    return send_file(csv_bytes, mimetype='text/csv; charset=utf-8',
                     as_attachment=True, download_name=f'{tipo}_{datetime.now().strftime("%Y%m%d")}.csv')

@api_rel_bp.route('/api/relatorio/pdf')
@login_required
def api_relatorio_pdf():
    tipo = request.args.get('tipo','vendas')
    dias = request.args.get('dias', 30, type=int)
    hoje = datetime.now().strftime('%d/%m/%Y')
    periodo = f'Período: {dias} dias até {hoje}'

    fn = PDF_FUNCTIONS.get(tipo)
    if not fn:
        return jsonify({'ok': False, 'erro': 'Tipo inválido'}), 400

    args = (dias,) if tipo in ('vendas','movimentacoes','lucro') else ()
    headers, data = fn(*args)

    titulos = {'vendas':'Vendas','stock':'Stock','produtos':'Produtos','movimentacoes':'Movimentações',
               'lucro':'Lucro','clientes':'Clientes','categoria':'Categoria'}
    pdf = gerar_pdf_relatorio(f'Relatório de {titulos.get(tipo,tipo)}', periodo, headers, data)
    return send_file(io.BytesIO(pdf), mimetype='application/pdf',
                     as_attachment=True, download_name=f'{tipo}_{datetime.now().strftime("%Y%m%d")}.pdf')

@api_rel_bp.route('/api/relatorio/xlsx')
@login_required
def api_relatorio_xlsx():
    tipo = request.args.get('tipo','vendas')
    dias = request.args.get('dias', 30, type=int)
    fn = PDF_FUNCTIONS.get(tipo)
    if not fn:
        return jsonify({'ok': False, 'erro': 'Tipo inválido'}), 400
    args = (dias,) if tipo in ('vendas','movimentacoes','lucro') else ()
    headers, data = fn(*args)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook()
    ws = wb.active

    titulos = {'vendas':'Vendas','stock':'Stock','produtos':'Produtos','movimentacoes':'Movimentações',
               'lucro':'Lucro','clientes':'Clientes','categoria':'Categoria'}
    ws.title = titulos.get(tipo, tipo)[:31]

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='1F6FEB')
    thin = Side(style='thin', color='D0D5DD')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(headers)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border

    for row in data:
        ws.append([str(c) if not isinstance(c, (int, float)) else c for c in row])

    for col in ws.columns:
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 45)
    ws.freeze_panes = 'A2'
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.border = border

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True,
                     download_name=f'{tipo}_{datetime.now().strftime("%Y%m%d")}.xlsx')
